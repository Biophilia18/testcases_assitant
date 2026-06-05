from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.api_models import API_COLUMN_TO_FIELD, API_EXCEL_COLUMNS, ApiTestCase
from src.text_utils import normalize_steps


def build_api_excel(api_cases: list[ApiTestCase], coverage_types: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "接口测试用例"
    sheet.append(API_EXCEL_COLUMNS)

    for case in api_cases:
        sheet.append([_cell_value(case, column) for column in API_EXCEL_COLUMNS])

    _style_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_value(case: ApiTestCase, column: str) -> str:
    field_name = API_COLUMN_TO_FIELD[column]
    value = getattr(case, field_name)
    if field_name == "steps":
        return normalize_steps(value)
    return value


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_alignment

    width_by_column = {
        "用例编号": 14,
        "模块": 16,
        "接口名称": 24,
        "请求方法": 12,
        "接口路径": 32,
        "请求参数": 30,
        "请求体": 36,
        "前置条件": 30,
        "操作步骤": 38,
        "预期状态码": 14,
        "预期结果": 42,
        "优先级": 10,
        "用例类型": 14,
        "备注": 28,
    }
    for index, column in enumerate(API_EXCEL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width_by_column.get(column, 18)

    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = _estimate_row_height(row)
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _estimate_row_height(row) -> int:
    max_lines = 1
    for cell in row:
        value = str(cell.value or "")
        max_lines = max(max_lines, value.count("\n") + 1)
    return min(max(36, max_lines * 22), 120)
