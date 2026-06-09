from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.api.models import API_COLUMN_TO_FIELD, API_EXCEL_COLUMNS, ApiDocument, ApiTestCase
from src.api.quality_checker import analyze_api_quality, api_quality_issues_to_rows, api_quality_summary
from src.text_utils import normalize_steps


def build_api_excel(
    api_cases: list[ApiTestCase],
    document: ApiDocument | None = None,
    coverage_types: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "接口测试用例"
    sheet.append(API_EXCEL_COLUMNS)

    for case in api_cases:
        sheet.append([_cell_value(case, column) for column in API_EXCEL_COLUMNS])

    _style_sheet(sheet)
    _write_quality_sheet(workbook, api_cases, document or ApiDocument())

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_value(case: ApiTestCase, column: str) -> str:
    field_name = API_COLUMN_TO_FIELD[column]
    value = getattr(case, field_name)
    if field_name == "steps":
        return normalize_steps(value)
    return value


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_alignment

    width_by_column = {
        "用例编号": 14,
        "用例标题": 30,
        "模块": 16,
        "接口名称": 24,
        "请求方法": 12,
        "接口路径": 32,
        "请求头": 30,
        "请求参数": 30,
        "请求体": 36,
        "前置条件": 30,
        "操作步骤": 38,
        "预期状态码": 14,
        "断言点": 42,
        "数据库校验": 34,
        "变量提取": 24,
        "优先级": 10,
        "用例类型": 14,
        "备注": 28,
    }
    for index, column in enumerate(API_EXCEL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width_by_column.get(column, 18)

    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = _estimate_row_height(row)
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_quality_sheet(workbook: Workbook, api_cases: list[ApiTestCase], document: ApiDocument) -> None:
    sheet = workbook.create_sheet("接口质量报告")
    issues = analyze_api_quality(api_cases, document)
    summary = api_quality_summary(api_cases, issues)

    sheet.append(["接口质量概览", ""])
    sheet.append(["项目/系统名称", document.project_name])
    sheet.append(["业务模块", document.module])
    sheet.append(["接口名称", document.api_name])
    sheet.append(["请求方法", document.method])
    sheet.append(["接口路径", document.path])
    sheet.append(["用例数", summary["case_count"]])
    sheet.append(["覆盖类型数", summary["covered_categories"]])
    sheet.append(["错误数", summary["error_count"]])
    sheet.append(["警告数", summary["warning_count"]])
    sheet.append([])
    sheet.append(["接口质量提示", "", "", "", ""])
    sheet.append(["分类", "级别", "涉及用例", "问题", "建议"])

    issue_rows = api_quality_issues_to_rows(issues)
    if issue_rows:
        for row in issue_rows:
            sheet.append([row["分类"], row["级别"], row["涉及用例"], row["问题"], row["建议"]])
    else:
        sheet.append(["-", "通过", "-", "暂无明显质量提示", "导出后仍建议测试人员结合接口文档和业务规则复核。"])

    _style_quality_sheet(sheet)


def _style_quality_sheet(sheet) -> None:
    title_fill = PatternFill("solid", fgColor="70AD47")
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    title_font = Font(color="FFFFFF", bold=True)
    header_font = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    for row_index in [1, 12]:
        for cell in sheet[row_index]:
            cell.fill = title_fill
            cell.font = title_font
            cell.border = thin_border
            cell.alignment = wrap_alignment

    for cell in sheet[13]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = wrap_alignment

    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = _estimate_row_height(row)
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment

    widths = [18, 16, 22, 42, 52]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A14"


def _estimate_row_height(row) -> int:
    max_lines = 1
    for cell in row:
        value = str(cell.value or "")
        max_lines = max(max_lines, value.count("\n") + 1)
    return min(max(36, max_lines * 22), 120)
