from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.models import COLUMN_TO_FIELD, EXCEL_COLUMNS, TestCase
from src.quality_checker import build_quality_report_rows
from src.requirement_trace import build_requirement_trace, requirement_trace_to_rows
from src.text_utils import normalize_steps


def build_excel(cases: list[TestCase], coverage_types: list[str] | None = None, requirement_text: str = "") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试用例"

    sheet.append(EXCEL_COLUMNS)

    for case in cases:
        sheet.append([_cell_value(case, column) for column in EXCEL_COLUMNS])

    _style_sheet(sheet)
    _build_quality_sheet(workbook, cases, coverage_types, requirement_text)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_alignment

    width_by_column = {
        "用例编号": 14,
        "模块": 16,
        "功能点": 20,
        "用例标题": 30,
        "优先级": 10,
        "前置条件": 36,
        "测试数据": 28,
        "操作步骤": 42,
        "预期结果": 42,
        "用例类型": 14,
        "备注": 28,
    }
    for index, column in enumerate(EXCEL_COLUMNS, start=1):
        width = width_by_column.get(column, 18)
        sheet.column_dimensions[get_column_letter(index)].width = width

    priority_fill = {
        "P0": PatternFill("solid", fgColor="F4CCCC"),
        "P1": PatternFill("solid", fgColor="FCE4D6"),
        "P2": PatternFill("solid", fgColor="FFF2CC"),
        "P3": PatternFill("solid", fgColor="E2F0D9"),
    }
    priority_col = EXCEL_COLUMNS.index("优先级") + 1

    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = _estimate_row_height(row)
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment
        row[priority_col - 1].alignment = center_alignment
        row[priority_col - 1].fill = priority_fill.get(str(row[priority_col - 1].value), PatternFill())

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _cell_value(case: TestCase, column: str) -> str:
    field_name = COLUMN_TO_FIELD[column]
    value = getattr(case, field_name)
    if field_name == "steps":
        return normalize_steps(value)
    return value


def _estimate_row_height(row) -> int:
    max_lines = 1
    for cell in row:
        value = str(cell.value or "")
        max_lines = max(max_lines, value.count("\n") + 1)
    return min(max(36, max_lines * 22), 120)


def _build_quality_sheet(
    workbook: Workbook,
    cases: list[TestCase],
    coverage_types: list[str] | None = None,
    requirement_text: str = "",
) -> None:
    sheet = workbook.create_sheet("质量报告")
    for row in build_quality_report_rows(cases, coverage_types):
        sheet.append(row)

    trace_items = build_requirement_trace(cases, requirement_text)
    if trace_items:
        sheet.append([])
        sheet.append(["需求规则覆盖追踪"])
        sheet.append(["来源", "规则内容", "是否覆盖", "命中用例", "建议"])
        for row in requirement_trace_to_rows(trace_items):
            sheet.append([row["来源"], row["规则内容"], row["是否覆盖"], row["命中用例"], row["建议"]])

    header_fill = PatternFill("solid", fgColor="70AD47")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        if row[0].value in {"需求规则覆盖追踪", "来源"}:
            for cell in row:
                cell.fill = header_fill
                cell.font = header_font

    for column_index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 22 if column_index < 3 else 50

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
