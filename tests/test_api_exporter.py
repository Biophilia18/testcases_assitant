from io import BytesIO

from openpyxl import load_workbook

from src.api_exporter import build_api_excel
from src.api_models import API_EXCEL_COLUMNS, ApiDocument, ApiTestCase


def test_build_api_excel_contains_api_sheet_and_data():
    case = ApiTestCase(
        case_id="API-01-01",
        module="设备控制",
        api_name="设备控制接口",
        method="POST",
        path="/api/devices/{deviceId}/control",
        headers="Authorization: Bearer token",
        query_params="deviceId=10001",
        request_body='{"action": "open"}',
        precondition="已获取有效 token",
        steps="1. 构造请求\n2. 发送请求",
        expected_status="200",
        assertions="接口返回成功",
        db_check="设备状态记录更新",
        extract_vars="",
        priority="P1",
        case_type="接口测试",
        remark="示例",
    )

    data = build_api_excel([case])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["接口测试用例"]

    headers = [cell.value for cell in sheet[1]]
    row = [cell.value for cell in sheet[2]]

    assert headers == API_EXCEL_COLUMNS
    assert row[0] == "API-01-01"
    assert row[2] == "设备控制接口"
    assert row[3] == "POST"
    assert row[5] == "Authorization: Bearer token"
    assert row[10] == "200"
    assert row[11] == "接口返回成功"
    assert row[12] == "设备状态记录更新"


def test_build_api_excel_contains_quality_report_sheet():
    case = ApiTestCase(
        case_id="API-01-01",
        module="设备控制",
        api_name="设备控制接口",
        method="POST",
        path="/api/devices/{deviceId}/control",
        headers="Authorization: Bearer token",
        query_params="deviceId=10001",
        request_body='{"action": "open"}',
        precondition="已获取有效 token",
        steps="1. 构造请求\n2. 发送请求",
        expected_status="200",
        assertions="成功",
        db_check="",
        extract_vars="",
        priority="P1",
        case_type="正常请求",
        remark="示例",
    )
    document = ApiDocument(
        project_name="智控家监测系统",
        module="设备控制",
        api_name="设备控制接口",
        method="POST",
        path="/api/devices/{deviceId}/control",
        auth="Bearer Token",
        response_example='{"code":0}',
        db_checks="设备状态记录更新",
    )

    data = build_api_excel([case], document=document)
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["接口质量报告"]

    assert "接口测试用例" in workbook.sheetnames
    assert "接口质量报告" in workbook.sheetnames
    assert sheet["A1"].value == "接口质量概览"
    assert sheet["A12"].value == "接口质量提示"
    assert sheet["A13"].value == "分类"
    assert any(row[0] == "数据库校验" for row in sheet.iter_rows(min_row=14, values_only=True))
