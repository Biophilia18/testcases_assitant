from io import BytesIO

from openpyxl import load_workbook

from src.exporter import build_excel
from src.models import EXCEL_COLUMNS, TestCase


def test_build_excel_contains_new_columns_and_data():
    case = TestCase(
        case_id="TC-01-01",
        module="订单",
        feature="查询订单",
        title="查询订单-正常流程",
        precondition="用户已登录",
        test_data="订单号：A001",
        steps="1. 输入订单号\n2. 查询",
        expected_result="展示订单详情",
        priority="P1",
        case_type="功能测试",
        remark="示例",
    )

    data = build_excel([case])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["测试用例"]

    headers = [cell.value for cell in sheet[1]]
    row = [cell.value for cell in sheet[2]]

    assert headers == EXCEL_COLUMNS
    assert row[0] == "TC-01-01"
    assert row[5] == "订单号：A001"
    assert row[10] == "示例"
    assert row[2] == "查询订单-正常流程"
    assert row[9] == "查询订单"


def test_build_excel_normalizes_steps_and_sets_row_height():
    case = TestCase(
        case_id="TC-01-01",
        module="订单",
        feature="查询订单",
        title="查询订单-正常流程",
        precondition="用户已登录",
        test_data="订单号：A001",
        steps="1. 输入订单号 2. 点击查询 3. 查看结果",
        expected_result="展示订单详情",
        priority="P1",
        case_type="功能测试",
        remark="示例",
    )

    data = build_excel([case])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["测试用例"]

    assert sheet["G2"].value == "1. 输入订单号\n2. 点击查询\n3. 查看结果"
    assert sheet.row_dimensions[2].height >= 66


def test_build_excel_styles_header_and_priority():
    case = TestCase(
        case_id="TC-01-01",
        module="订单",
        feature="查询订单",
        title="查询订单-正常流程",
        precondition="用户已登录",
        test_data="订单号：A001",
        steps="1. 查询\n2. 查看",
        expected_result="展示订单详情",
        priority="P1",
        case_type="功能测试",
        remark="示例",
    )

    data = build_excel([case])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["测试用例"]

    assert sheet["A1"].fill.fgColor.rgb == "001F4E78"
    assert sheet["D2"].fill.fgColor.rgb == "00FCE4D6"


def test_build_excel_adds_quality_report_sheet():
    case = TestCase(
        case_id="TC-01-01",
        module="订单",
        feature="查询订单",
        title="查询订单-正常流程",
        precondition="用户已登录",
        test_data="订单号：A001",
        steps="1. 查询\n2. 查看",
        expected_result="展示订单详情",
        priority="P1",
        case_type="功能测试",
        remark="示例",
    )

    data = build_excel([case])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["质量报告"]

    assert "质量报告" in workbook.sheetnames
    assert sheet["A1"].value == "指标"
    assert sheet["A2"].value == "总用例数"
    assert sheet["B2"].value == 1


def test_build_excel_quality_report_records_coverage_types():
    case = TestCase(
        case_id="TC-01-01",
        module="订单",
        feature="查询订单",
        title="查询订单-正常流程",
        precondition="用户已登录",
        test_data="订单号：A001",
        steps="1. 查询\n2. 查看",
        expected_result="展示订单详情",
        priority="P1",
        case_type="功能测试",
        remark="覆盖类型：正常流程",
    )

    data = build_excel([case], coverage_types=["正常流程", "弱网/超时"])
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["质量报告"]

    assert sheet["A5"].value == "覆盖类型"
    assert sheet["B5"].value == "正常流程、弱网/超时"
